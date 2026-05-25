from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass
class ParentRef:
    value: str
    cell: str
    sheet: str = "oligomers"


@dataclass
class RegenSheet:
    subdir: Path
    xlsx_path: Path
    identity: str
    parents: list[ParentRef] = field(default_factory=list)
    is_root: bool = False
    workbook_handle: Any = None

    def load_workbook(self):
        if self.workbook_handle is None:
            self.workbook_handle = load_workbook(self.xlsx_path, data_only=False)
        return self.workbook_handle


def find_xlsx_in_subdir(subdir: Path) -> list[Path]:
    """Find all Data Entry_*.xlsx files in a subdirectory."""
    return sorted(subdir.glob("Data Entry_*.xlsx"))


def parse_regen_sheet(subdir: Path, resolver: Any = None) -> RegenSheet | None:
    """
    Parse a single experiment subdirectory.
    Returns RegenSheet with parents extracted from oligomers tab, or None if invalid.
    """
    xlsx_files = find_xlsx_in_subdir(subdir)

    if len(xlsx_files) == 0:
        return None
    if len(xlsx_files) > 1:
        return None

    xlsx_path = xlsx_files[0]

    try:
        wb = load_workbook(xlsx_path, data_only=False, read_only=True)
    except Exception:
        return None

    if "oligomers" not in wb.sheetnames:
        identity = subdir.name
        return RegenSheet(
            subdir=subdir, xlsx_path=xlsx_path, identity=identity, parents=[], is_root=True
        )

    ws = wb["oligomers"]
    parents = []

    for row_idx in range(2, ws.max_row + 1):
        cell_a = ws[f"A{row_idx}"]

        if cell_a.value is None or str(cell_a.value).strip() == "":
            break

        cell_text = str(cell_a.value).strip()

        if cell_text.upper() == "PROCEDURE":
            break

        parents.append(ParentRef(value=cell_text, cell=f"A{row_idx}", sheet="oligomers"))

    identity = subdir.name
    is_root = len(parents) == 0

    sheet = RegenSheet(
        subdir=subdir, xlsx_path=xlsx_path, identity=identity, parents=parents, is_root=is_root
    )
    sheet.workbook_handle = wb
    return sheet
