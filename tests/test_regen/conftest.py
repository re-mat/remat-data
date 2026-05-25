from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook


def create_regen_xlsx(path: Path, parent_refs: list[str] | None = None):
    """
    Create a minimal Data Entry_*.xlsx with an oligomers tab.

    Args:
        path: Where to save the xlsx
        parent_refs: List of parent reference strings to add to oligomers!A2:A(n)
    """
    if parent_refs is None:
        parent_refs = []

    wb = Workbook()
    wb.remove(wb.active)

    ws_oligo = wb.create_sheet("oligomers")
    ws_oligo["A1"] = "Oligo ID"
    ws_oligo["B1"] = "Measured mass (g)"

    for idx, parent_ref in enumerate(parent_refs, start=2):
        ws_oligo[f"A{idx}"] = parent_ref
        ws_oligo[f"B{idx}"] = 2.0

    ws_oligo["A13"] = "PROCEDURE"

    ws_general = wb.create_sheet("general")
    ws_general["A1"] = "Operator Initials"
    ws_general["B1"] = "AR"

    wb.save(path)


@pytest.fixture
def temp_submission_dir() -> Path:
    """Create a temporary directory for testing submissions."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def simple_chain(temp_submission_dir: Path) -> Path:
    """Create Gen 0 -> Gen 1 -> Gen 2 chain."""
    for gen_name in ["Gen 0", "Gen 1", "Gen 2"]:
        subdir = temp_submission_dir / gen_name
        subdir.mkdir(parents=True)
        xlsx_path = subdir / "Data Entry_Gen0.xlsx"
        if gen_name == "Gen 0":
            create_regen_xlsx(xlsx_path, parent_refs=[])
        elif gen_name == "Gen 1":
            create_regen_xlsx(xlsx_path, parent_refs=["Gen 0"])
        elif gen_name == "Gen 2":
            create_regen_xlsx(xlsx_path, parent_refs=["Gen 1"])

    return temp_submission_dir


@pytest.fixture
def diamond(temp_submission_dir: Path) -> Path:
    """Create diamond graph: A -> B, A -> C, {B,C} -> D."""
    for name in ["A", "B", "C", "D"]:
        subdir = temp_submission_dir / name
        subdir.mkdir(parents=True)
        xlsx_path = subdir / "Data Entry_Test.xlsx"

        if name == "A":
            create_regen_xlsx(xlsx_path, parent_refs=[])
        elif name in ("B", "C"):
            create_regen_xlsx(xlsx_path, parent_refs=["A"])
        elif name == "D":
            create_regen_xlsx(xlsx_path, parent_refs=["B", "C"])

    return temp_submission_dir


@pytest.fixture
def multi_parent(temp_submission_dir: Path) -> Path:
    """Create fixture where one child has multiple parents."""
    for name in ["Parent1", "Parent2", "Child"]:
        subdir = temp_submission_dir / name
        subdir.mkdir(parents=True)
        xlsx_path = subdir / "Data Entry_Test.xlsx"

        if name.startswith("Parent"):
            create_regen_xlsx(xlsx_path, parent_refs=[])
        else:
            create_regen_xlsx(xlsx_path, parent_refs=["Parent1", "Parent2"])

    return temp_submission_dir


@pytest.fixture
def multi_child(temp_submission_dir: Path) -> Path:
    """Create fixture where one parent has multiple children."""
    for name in ["Parent", "Child1", "Child2"]:
        subdir = temp_submission_dir / name
        subdir.mkdir(parents=True)
        xlsx_path = subdir / "Data Entry_Test.xlsx"

        if name == "Parent":
            create_regen_xlsx(xlsx_path, parent_refs=[])
        else:
            create_regen_xlsx(xlsx_path, parent_refs=["Parent"])

    return temp_submission_dir


@pytest.fixture
def roots_only(temp_submission_dir: Path) -> Path:
    """Create 3 independent root experiments."""
    for name in ["Exp1", "Exp2", "Exp3"]:
        subdir = temp_submission_dir / name
        subdir.mkdir(parents=True)
        xlsx_path = subdir / "Data Entry_Test.xlsx"
        create_regen_xlsx(xlsx_path, parent_refs=[])

    return temp_submission_dir


@pytest.fixture
def missing_parent(temp_submission_dir: Path) -> Path:
    """Create fixture where child references non-existent parent."""
    subdir = temp_submission_dir / "Child"
    subdir.mkdir(parents=True)
    xlsx_path = subdir / "Data Entry_Test.xlsx"
    create_regen_xlsx(xlsx_path, parent_refs=["NonExistentParent"])

    return temp_submission_dir


@pytest.fixture
def cycle_fixture(temp_submission_dir: Path) -> Path:
    """Create cycle: A <-> B."""
    for name in ["A", "B"]:
        subdir = temp_submission_dir / name
        subdir.mkdir(parents=True)

    create_regen_xlsx(
        temp_submission_dir / "A" / "Data Entry_Test.xlsx", parent_refs=["B"]
    )
    create_regen_xlsx(
        temp_submission_dir / "B" / "Data Entry_Test.xlsx", parent_refs=["A"]
    )

    return temp_submission_dir


@pytest.fixture
def self_reference(temp_submission_dir: Path) -> Path:
    """Create fixture where an experiment references itself."""
    subdir = temp_submission_dir / "SelfRef"
    subdir.mkdir(parents=True)
    xlsx_path = subdir / "Data Entry_Test.xlsx"
    create_regen_xlsx(xlsx_path, parent_refs=["SelfRef"])

    return temp_submission_dir


@pytest.fixture
def no_xlsx(temp_submission_dir: Path) -> Path:
    """Create subdir with no xlsx file."""
    subdir = temp_submission_dir / "NoXlsx"
    subdir.mkdir(parents=True)
    (subdir / "dummy.txt").write_text("dummy")

    return temp_submission_dir


@pytest.fixture
def multiple_xlsx(temp_submission_dir: Path) -> Path:
    """Create subdir with multiple xlsx files."""
    subdir = temp_submission_dir / "MultiXlsx"
    subdir.mkdir(parents=True)

    create_regen_xlsx(subdir / "Data Entry_1.xlsx", parent_refs=[])
    create_regen_xlsx(subdir / "Data Entry_2.xlsx", parent_refs=[])

    return temp_submission_dir
