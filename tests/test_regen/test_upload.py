from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

from openpyxl import load_workbook

from remat_data.regen.upload import _add_parent_url_column, upload_regen_directory
from remat_data.regen.validate import validate_directory

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_mock_clowder(ids: list[str] | None = None):
    """Return a MagicMock ClowderClient whose post() returns sequential dataset IDs."""
    clowder = MagicMock()
    id_iter = iter(ids or [])

    def _post(endpoint, _payload=None):
        if endpoint == "/datasets/createempty":
            return {"id": next(id_iter)}
        return {}

    clowder.post.side_effect = _post
    clowder.post_file.return_value = "file_id_stub"
    return clowder


# ---------------------------------------------------------------------------
# test_add_parent_url_column
# ---------------------------------------------------------------------------


class TestAddParentUrlColumn:
    def test_appends_url_column(self, simple_chain: Path):
        report = validate_directory(simple_chain)
        assert report.ok

        sheet_gen1 = report.sheets["Gen 1"]
        dataset_id_map = {"Gen 0": "abc123"}

        _add_parent_url_column(sheet_gen1, dataset_id_map, "space99")

        wb = load_workbook(sheet_gen1.xlsx_path, data_only=False)
        ws = wb["oligomers"]

        # Header is always in column C (index 3)
        assert ws.cell(row=1, column=3).value == "Parent Dataset URL"

        # URL in column C of the parent ref row
        url_cell = ws.cell(row=2, column=3).value
        assert url_cell is not None
        assert "abc123" in url_cell
        assert "space99" in url_cell

    def test_col_a_unchanged(self, simple_chain: Path):
        report = validate_directory(simple_chain)
        sheet_gen1 = report.sheets["Gen 1"]
        dataset_id_map = {"Gen 0": "abc123"}

        _add_parent_url_column(sheet_gen1, dataset_id_map, "space99")

        wb = load_workbook(sheet_gen1.xlsx_path, data_only=False)
        ws = wb["oligomers"]
        assert ws["A2"].value == "Gen 0"

    def test_multiple_parents(self, multi_parent: Path):
        report = validate_directory(multi_parent)
        assert report.ok

        sheet_child = report.sheets["Child"]
        dataset_id_map = {"Parent1": "id_p1", "Parent2": "id_p2"}

        _add_parent_url_column(sheet_child, dataset_id_map, "spc")

        wb = load_workbook(sheet_child.xlsx_path, data_only=False)
        ws = wb["oligomers"]

        urls = [ws.cell(row=r, column=3).value for r in range(2, 4)]
        assert any("id_p1" in u for u in urls if u)
        assert any("id_p2" in u for u in urls if u)


# ---------------------------------------------------------------------------
# test_upload_regen_directory
# ---------------------------------------------------------------------------


class TestUploadRegenDirectory:
    def test_upload_order_simple_chain(self, simple_chain: Path):
        clowder = _make_mock_clowder(["id_gen0", "id_gen1", "id_gen2"])

        result = upload_regen_directory(simple_chain, clowder, "space_x")

        assert result == {"Gen 0": "id_gen0", "Gen 1": "id_gen1", "Gen 2": "id_gen2"}

    def test_parent_url_written_before_child_upload(self, simple_chain: Path):
        clowder = _make_mock_clowder(["id_gen0", "id_gen1", "id_gen2"])
        report_before = validate_directory(simple_chain)

        upload_regen_directory(simple_chain, clowder, "space_x")

        # After upload, Gen 1 xlsx should have the parent URL column
        wb = load_workbook(report_before.sheets["Gen 1"].xlsx_path)
        ws = wb["oligomers"]
        new_col = ws.max_column
        assert ws.cell(row=1, column=new_col).value == "Parent Dataset URL"
        assert "id_gen0" in (ws.cell(row=2, column=new_col).value or "")

    def test_roots_have_no_url_column_added(self, simple_chain: Path):
        clowder = _make_mock_clowder(["id_gen0", "id_gen1", "id_gen2"])
        report_before = validate_directory(simple_chain)

        upload_regen_directory(simple_chain, clowder, "space_x")

        wb = load_workbook(report_before.sheets["Gen 0"].xlsx_path)
        ws = wb["oligomers"]
        # Gen 0 is a root — column C should have no "Parent Dataset URL" header
        assert ws.cell(row=1, column=3).value != "Parent Dataset URL"

    def test_aborts_on_validation_failure(self, missing_parent: Path):
        clowder = _make_mock_clowder()

        result = upload_regen_directory(missing_parent, clowder, "space_x")

        assert result == {}
        clowder.post.assert_not_called()

    def test_all_experiments_uploaded(self, diamond: Path):
        clowder = _make_mock_clowder(["id_a", "id_b", "id_c", "id_d"])

        result = upload_regen_directory(diamond, clowder, "space_x")

        assert set(result.keys()) == {"A", "B", "C", "D"}
        # All 4 createempty calls made
        assert clowder.post.call_count == 4
